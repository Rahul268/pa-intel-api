"""
io_utils.py
File I/O utilities: read the submission manifest from Excel, list PDFs,
write result.csv, write JSONL audit records, and configure logging.
"""

from __future__ import annotations
import csv
import json
import logging
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

log = logging.getLogger(__name__)

# Exact column order required in result.csv
RESULT_COLUMNS = [
    "Filename",
    "Brand",
    "Age",
    "Step Therapy Requirements Documented in Policy",
    "Number of Steps through Brands",
    "Number of Steps through Generic",
    "Step through-Phototherapy",
    "TB Test required",
    "Quantity Limits",
    "Specialist Types",
    "Initial Authorization Duration(in-months)",
    "Reauthorization Duration(in-months)",
    "Reauthorization Required",
    "Reauthorization Requirements Documented in Policy",
    "Access Score",
]


# ── Logging setup ─────────────────────────────────────────────────────────────

def configure_logging(logs_dir: Path, level: int = logging.INFO) -> None:
    """Set up root logger to write to both console and logs/pipeline.log."""
    logs_dir.mkdir(parents=True, exist_ok=True)
    log_file = logs_dir / "pipeline.log"

    fmt = "%(asctime)s | %(levelname)-8s | %(name)s | %(message)s"
    datefmt = "%Y-%m-%d %H:%M:%S"

    handlers: list[logging.Handler] = [
        logging.StreamHandler(sys.stdout),
        logging.FileHandler(log_file, encoding="utf-8"),
    ]
    logging.basicConfig(level=level, format=fmt, datefmt=datefmt, handlers=handlers)
    log.info("Logging initialised → %s", log_file)


# ── Manifest ──────────────────────────────────────────────────────────────────

def read_submissions_manifest(
    xlsx_path: Path, sheet_name: str = "Submissions"
) -> list[dict[str, str]]:
    """
    Read the Submissions tab and return a deduplicated list of
    {Filename, Brand} dicts in original row order.
    """
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Business rules workbook not found: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' not found in {xlsx_path.name}. "
            f"Available sheets: {wb.sheetnames}"
        )

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Sheet '{sheet_name}' is empty.")

    # Locate Filename and Brand columns by header name (row 0)
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    try:
        fn_idx = headers.index("Filename")
        brand_idx = headers.index("Brand")
    except ValueError as exc:
        raise ValueError(
            f"Expected 'Filename' and 'Brand' columns in sheet '{sheet_name}'. "
            f"Found: {headers}"
        ) from exc

    seen: set[tuple[str, str]] = set()
    manifest: list[dict[str, str]] = []
    for row in rows[1:]:
        fn = row[fn_idx]
        brand = row[brand_idx]
        if fn is None or brand is None:
            continue
        fn = str(fn).strip()
        brand = str(brand).strip()
        if not fn or not brand:
            continue
        key = (fn, brand)
        if key not in seen:
            seen.add(key)
            manifest.append({"Filename": fn, "Brand": brand})

    log.info("Manifest loaded: %d Filename+Brand rows", len(manifest))
    return manifest


# ── PDF listing ───────────────────────────────────────────────────────────────

def list_pdfs(directory: Path) -> list[Path]:
    """Return sorted list of PDF paths in a directory."""
    if not directory.exists():
        return []
    pdfs = sorted(directory.glob("*.pdf"))
    log.debug("Found %d PDFs in %s", len(pdfs), directory)
    return pdfs


# ── Output writers ────────────────────────────────────────────────────────────

def ensure_dirs(*paths: Path) -> None:
    """Create all given directories (and parents) if they do not exist."""
    for p in paths:
        p.mkdir(parents=True, exist_ok=True)


def write_result_csv(rows: list[dict[str, Any]], output_path: Path) -> None:
    """Write result.csv with exact column order. Fills missing fields with 'NA'."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=RESULT_COLUMNS,
            extrasaction="ignore",
        )
        writer.writeheader()
        for row in rows:
            clean = {col: row.get(col, "NA") or "NA" for col in RESULT_COLUMNS}
            writer.writerow(clean)
    log.info("Result CSV written: %s (%d rows)", output_path, len(rows))


def write_result_excel(rows: list[dict[str, Any]], output_path: Path) -> None:
    """
    Write result.xlsx with a 'Submissions' sheet whose column order, names,
    and data exactly match the CSV produced by write_result_csv.
    Columns follow RESULT_COLUMNS — the same order as the Submissions tab in
    PA_Business_Rules.xlsx.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        log.warning("openpyxl not installed — Excel output skipped. Run: pip install openpyxl")
        return

    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Submissions"

    header_font  = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    header_fill  = PatternFill(fill_type="solid", fgColor="2E4057")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin         = Side(style="thin")
    border       = Border(left=thin, right=thin, top=thin, bottom=thin)
    data_font    = Font(name="Calibri", size=10)
    alt_fill     = PatternFill(fill_type="solid", fgColor="F2F6FA")
    wrap_cols    = {
        "Step Therapy Requirements Documented in Policy",
        "Reauthorization Requirements Documented in Policy",
    }
    col_widths   = {
        "Filename": 28, "Brand": 14, "Age": 10,
        "Step Therapy Requirements Documented in Policy": 55,
        "Number of Steps through Brands": 14,
        "Number of Steps through Generic": 14,
        "Step through-Phototherapy": 14, "TB Test required": 12,
        "Quantity Limits": 30, "Specialist Types": 20,
        "Initial Authorization Duration(in-months)": 18,
        "Reauthorization Duration(in-months)": 18,
        "Reauthorization Required": 14,
        "Reauthorization Requirements Documented in Policy": 55,
        "Access Score": 12,
    }

    ws.row_dimensions[1].height = 40
    for ci, col in enumerate(RESULT_COLUMNS, 1):
        cell = ws.cell(row=1, column=ci, value=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
        ws.column_dimensions[get_column_letter(ci)].width = col_widths.get(col, 15)

    for ri, row in enumerate(rows, 2):
        fill = alt_fill if ri % 2 == 0 else PatternFill()
        for ci, col in enumerate(RESULT_COLUMNS, 1):
            val  = row.get(col, "NA") or "NA"
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font      = data_font
            cell.fill      = fill
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=(col in wrap_cols))
            cell.border    = border

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(RESULT_COLUMNS))}1"

    wb.save(output_path)
    log.info("Result Excel written: %s (%d rows)", output_path, len(rows))


def write_jsonl(records: list[dict[str, Any]], output_path: Path) -> None:
    """Append records to a JSONL file (one JSON object per line)."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "a", encoding="utf-8") as f:
        for rec in records:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")


def timestamped_backup(path: Path) -> Path | None:
    """If path exists, rename it with a timestamp suffix and return the new path."""
    if path.exists():
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = path.with_name(f"{path.stem}_{ts}{path.suffix}")
        path.rename(backup)
        log.info("Backed up existing file: %s → %s", path.name, backup.name)
        return backup
    return None
